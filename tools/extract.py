# -*- coding: utf-8 -*-
"""엑셀 가계부 -> JSON 추출 + 정산 로직 검증.

읽기 전용이다. 엑셀을 수정하지 않는다(쓰기는 반드시 Excel COM, CLAUDE.md 참조).

출력 두 가지:
  data/ledger.json          검증용 (엑셀 계산값 포함). **git에 올리지 않는다**
  data/가계부_복원.json      앱의 "백업 파일에서 복원"으로 폰에 넣는 파일

둘 다 실제 가계부 내역이라 저장소 밖(`data/`, .gitignore 처리)에 둔다.
GitHub Pages 무료 등급은 공개 저장소만 지원하므로, 소스에 데이터를 두면
그대로 인터넷에 공개된다.
"""
import json
import datetime as dt
from pathlib import Path

import openpyxl

SRC_DIR = Path(r"C:\Users\howra\Desktop\회계마감")
DATA_DIR = Path(__file__).resolve().parent.parent / "data"
OUT = DATA_DIR / "ledger.json"
OUT_BACKUP = DATA_DIR / "가계부_복원.json"

YEARS = [2024, 2025, 2026]

# 세대별 입력 영역: (세대키, 일자열, 범주열, 상세열, 결제수단열, 금액열)
HOUSEHOLDS = [
    ("eunji", 2, 3, 4, 5, 6),    # 은지네 B~F
    ("jisu", 7, 8, 9, 10, 11),   # 지수네 G~K
]
ROW_START, ROW_END = 20, 52


def parse_day(value):
    """일자 셀을 1~31 정수로. '19일' 문자열과 datetime이 섞여 있다."""
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.day
    s = str(value).strip()
    if not s:
        return None
    digits = "".join(ch for ch in s if ch.isdigit())
    if not digits:
        return None
    day = int(digits)
    return day if 1 <= day <= 31 else None


def extract_year(year):
    path = SRC_DIR / f"{year}년_가계부.xlsx"
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)

    txns, months = [], []

    for m in range(1, 13):
        name = f"{m}월"
        if name not in wbf.sheetnames:
            continue
        sf, sv = wbf[name], wbv[name]

        # 이 시트가 2세대 정산 구조인지 판별.
        # 원본 템플릿은 N2='총 지출', 정산 구조는 N2='은지네'.
        is_split = str(sf["N2"].value or "").strip() == "은지네"

        last_day = None
        for hkey, c_day, c_cat, c_desc, c_pay, c_amt in HOUSEHOLDS:
            for r in range(ROW_START, ROW_END + 1):
                amount = sf.cell(r, c_amt).value
                if amount is None or not isinstance(amount, (int, float)):
                    continue

                day = parse_day(sf.cell(r, c_day).value)
                # 엑셀은 같은 날 연속 입력 시 일자를 비워둔다 -> 직전 값 승계
                if day is None:
                    day = last_day
                else:
                    last_day = day

                txns.append({
                    "year": year,
                    "month": m,
                    "day": day,
                    "household": hkey,
                    "category": str(sf.cell(r, c_cat).value or "").strip(),
                    "description": str(sf.cell(r, c_desc).value or "").strip(),
                    "payment": str(sf.cell(r, c_pay).value or "").strip(),
                    "amount": int(round(float(amount))),
                    "source": f"{year}/{name}/r{r}",
                })
            last_day = None  # 세대가 바뀌면 승계 초기화

        months.append({
            "year": year,
            "month": m,
            "isSplit": is_split,
            # 엑셀이 계산해둔 값 (검증 기준)
            "xl": {
                "eunjiPaid": sv["P2"].value,
                "jisuPaid": sv["P4"].value,
                "total": sv["P6"].value,
                "eunjiShare": sv["P8"].value,
                "jisuShare": sv["P10"].value,
                "eunjiGap": sv["P12"].value,
                "cumulative": sv["Q12"].value,
            },
            # 1월에만 존재하는 전년도 이월값
            "carryIn": sf["N16"].value if isinstance(sf["N16"].value, (int, float)) else None,
            "ratioFormula": {"eunji": str(sf["Q2"].value), "jisu": str(sf["Q4"].value)},
        })

    return txns, months


def main():
    all_txns, all_months = [], []
    for y in YEARS:
        t, m = extract_year(y)
        all_txns += t
        all_months += m

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps({"transactions": all_txns, "months": all_months},
                   ensure_ascii=False, indent=1),
        encoding="utf-8",
    )

    split = [m for m in all_months if m["isSplit"]]
    split_keys = {(m["year"], m["month"]) for m in split}

    # 앱이 바로 읽는 백업 형식. 2세대 정산이 시작된 2024/9 이후만 담는다.
    restore_txns = [
        {k: v for k, v in t.items() if k != "source"}
        for t in all_txns
        if (t["year"], t["month"]) in split_keys
    ]
    OUT_BACKUP.write_text(
        json.dumps(
            {
                "app": "mageum",
                "version": 1,
                "exportedAt": dt.datetime.now().isoformat(),
                "transactions": restore_txns,
                # 엑셀에서 수식 안에 섞여 있던 수동 보정 (CLAUDE.md 참조)
                "adjustments": [
                    {"year": 2024, "month": 9, "amount": -55644,
                     "reason": "기존 재산세 납부 제외 (정산 시작 시 보정)"},
                    {"year": 2026, "month": 5, "amount": 30000,
                     "reason": "엑셀 5월 수식에 있던 보정 (근거 미확인)"},
                ],
            },
            ensure_ascii=False, indent=1,
        ),
        encoding="utf-8",
    )

    lines = [f"거래 {len(all_txns)}건, 월 {len(all_months)}개 -> {OUT}"]
    lines.append(f"2세대 정산 구조 월: {len(split)}개")
    ratios = {(m["ratioFormula"]["eunji"], m["ratioFormula"]["jisu"]) for m in split}
    lines.append(f"비율 수식 종류: {ratios}")
    lines.append(f"복원용 {len(restore_txns)}건 -> {OUT_BACKUP}")
    lines.append("  이 파일을 폰으로 옮겨 앱 설정 > '백업 파일에서 복원' 으로 넣는다.")
    Path(__file__).parent.joinpath("extract_report.txt").write_text(
        "\n".join(lines), encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
